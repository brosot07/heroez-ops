import docx
import pytest

from heroez_ops import doctemplate as dtp
from heroez_ops.terbilang import parse_amount, rupiah, terbilang


def sample_contract() -> bytes:
    """A small 'approved' contract: split runs, bold party name, SOW table, header number."""
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "No. 001/HRZ-KTR/X/2026"
    d.add_heading("PERJANJIAN KERJA SAMA KAMPANYE", level=1)
    p = d.add_paragraph("Perjanjian ini dibuat antara PT Heroez Kreasi Indonesia dan ")
    r = p.add_run("PT Contoh ")
    r.bold = True
    r2 = p.add_run("Brand Tbk")
    r2.bold = True
    p.add_run(", beralamat di Jl. Sudirman No. 1, Jakarta.")
    d.add_paragraph("Pasal 1. Para Pihak sepakat untuk bekerja sama dengan itikad baik.")
    t = d.add_table(rows=3, cols=4)
    for c, h in zip(t.rows[0].cells, ["No", "Talent", "Platform", "Fee"]):
        c.text = h
    for c, v in zip(t.rows[1].cells, ["1", "Dewi Lestari", "TikTok", "Rp10.000.000"]):
        c.text = v
    for c, v in zip(t.rows[2].cells, ["2", "Raka", "Instagram", "Rp5.000.000"]):
        c.text = v
    d.add_paragraph("Nilai kontrak sebesar Rp15.000.000 (lima belas juta Rupiah) belum termasuk pajak.")
    d.add_paragraph("Ditandatangani pada 1 Oktober 2026.")
    buf = __import__("io").BytesIO()
    d.save(buf)
    return buf.getvalue()


def spec_for(doc_bytes: bytes) -> dict:
    blocks = {dtp.run_text(b.paragraph): b.id for b in dtp.blocks(dtp.load(doc_bytes))}
    party = next(i for t, i in blocks.items() if "PT Contoh" in t)
    fee = next(i for t, i in blocks.items() if "Nilai kontrak" in t)
    signed = next(i for t, i in blocks.items() if "Ditandatangani" in t)
    header = next(i for t, i in blocks.items() if t.startswith("No. "))
    return {
        "fields": [
            {"key": "client_name", "label": "Nama klien", "type": "text",
             "occurrences": [{"block": party, "text": "PT Contoh Brand Tbk"}]},
            {"key": "client_address", "label": "Alamat klien", "type": "text",
             "occurrences": [{"block": party, "text": "Jl. Sudirman No. 1, Jakarta"}]},
            {"key": "fee_total", "label": "Total fee", "type": "money",
             "occurrences": [{"block": fee, "text": "Rp15.000.000"}]},
            {"key": "fee_words", "label": "Terbilang", "type": "money_words", "source": "fee_total",
             "occurrences": [{"block": fee, "text": "lima belas juta Rupiah"}]},
            {"key": "sign_date", "label": "Tanggal", "type": "date",
             "occurrences": [{"block": signed, "text": "1 Oktober 2026"}]},
            {"key": "contract_number", "label": "Nomor", "type": "contract_number",
             "occurrences": [{"block": header, "text": "001/HRZ-KTR/X/2026"}]},
        ],
        "deliverables": {"table": 0, "template_row": 1, "sample_rows_after": 1, "columns": [
            {"col": 0, "key": "no", "label": "No"}, {"col": 1, "key": "talent", "label": "Talent"},
            {"col": 2, "key": "platform", "label": "Platform"}, {"col": 3, "key": "fee", "label": "Fee", "type": "money"}]},
    }


def test_mark_and_render_keeps_wording_and_formatting():
    original = sample_contract()
    spec = spec_for(original)
    template, problems = dtp.mark(original, spec)
    assert problems == []
    assert dtp.placeholders(template) >= {"client_name", "fee_total", "row.talent", "contract_number"}

    out = dtp.render(template, spec, {
        "client_name": "PT Glowskin Indonesia", "client_address": "Jl. Kemang Raya 5, Jakarta Selatan",
        "fee_total": 150_000_000, "sign_date": "2026-11-10", "contract_number": "014/HRZ-KTR/XI/2026",
    }, rows=[{"talent": "Dewi", "platform": "TikTok", "fee": 100_000_000},
             {"talent": "Raka", "platform": "IG Reels", "fee": 30_000_000},
             {"talent": "Sinta", "platform": "YouTube", "fee": 20_000_000}])
    d = dtp.load(out)
    body = "\n".join(p.text for p in d.paragraphs)
    assert "Perjanjian ini dibuat antara PT Heroez Kreasi Indonesia dan PT Glowskin Indonesia, beralamat di Jl. Kemang Raya 5, Jakarta Selatan." in body
    assert "Pasal 1. Para Pihak sepakat untuk bekerja sama dengan itikad baik." in body       # untouched clause
    assert "Rp150.000.000 (seratus lima puluh juta Rupiah)" in body
    assert "10 November 2026" in body
    assert d.sections[0].header.paragraphs[0].text == "No. 014/HRZ-KTR/XI/2026"
    # client name keeps the bold formatting it had in the approved template
    party = next(p for p in d.paragraphs if "Glowskin" in p.text)
    assert any(r.bold and "Glowskin" in r.text for r in party.runs)
    rows = [[c.text for c in r.cells] for r in d.tables[0].rows]
    assert rows == [["No", "Talent", "Platform", "Fee"], ["1", "Dewi", "TikTok", "Rp100.000.000"],
                    ["2", "Raka", "IG Reels", "Rp30.000.000"], ["3", "Sinta", "YouTube", "Rp20.000.000"]]
    assert "{{" not in body


def test_round_trip_with_original_values_reproduces_original():
    original = sample_contract()
    spec = spec_for(original)
    template, _ = dtp.mark(original, spec)
    out = dtp.render(template, spec, {"client_name": "PT Contoh Brand Tbk", "client_address": "Jl. Sudirman No. 1, Jakarta",
                                      "fee_total": 15_000_000, "sign_date": "2026-10-01",
                                      "contract_number": "001/HRZ-KTR/X/2026"},
                     rows=[{"talent": "Dewi Lestari", "platform": "TikTok", "fee": 10_000_000},
                           {"talent": "Raka", "platform": "Instagram", "fee": 5_000_000}])
    assert dtp.outline(dtp.load(out)) == dtp.outline(dtp.load(original))


def test_bad_occurrence_is_reported():
    original = sample_contract()
    spec = {"fields": [{"key": "x", "label": "X", "type": "text", "occurrences": [{"block": "B2", "text": "tidak ada"}]}]}
    _, problems = dtp.mark(original, spec)
    assert problems and "tidak ditemukan" in problems[0]


@pytest.mark.parametrize("n,words", [(0, "nol"), (11, "sebelas"), (15, "lima belas"), (100, "seratus"),
                                     (1000, "seribu"), (1_500_000, "satu juta lima ratus ribu"),
                                     (150_000_000, "seratus lima puluh juta"),
                                     (2_350_750_000, "dua miliar tiga ratus lima puluh juta tujuh ratus lima puluh ribu")])
def test_terbilang(n, words):
    assert terbilang(n) == words


def test_amount_parsing():
    assert parse_amount("Rp 150.000.000") == 150_000_000
    assert parse_amount("150jt") == 150_000_000
    assert parse_amount("1,5 M") == 1_500_000_000
    assert rupiah(1234567) == "Rp1.234.567"


def test_output_follows_template_style():
    fmt = dtp.infer_format("money", "Rp 15.000.000,-")
    assert dtp.format_value("money", 150_000_000, {}, fmt=fmt) == "Rp 150.000.000,-"
    fmt = dtp.infer_format("money", "IDR 15,000,000.00")
    assert dtp.format_value("money", 1_500_000, {}, fmt=fmt) == "IDR 1,500,000.00"
    w = dtp.infer_format("money_words", "LIMA BELAS JUTA RUPIAH")
    assert dtp.format_value("money_words", None, {"fee": 150_000_000}, "fee", w) == "SERATUS LIMA PULUH JUTA RUPIAH"
    w = dtp.infer_format("money_words", "Lima Belas Juta")
    assert dtp.format_value("money_words", None, {"fee": 2_000_000}, "fee", w) == "Dua Juta"
    assert dtp.format_value("date", "2026-11-05", {}, fmt=dtp.infer_format("date", "01/10/2026")) == "05/11/2026"
    assert dtp.format_value("date", "2026-11-05", {}, fmt=dtp.infer_format("date", "1 October 2026")) == "5 November 2026"
    assert dtp.format_value("date", "2026-11-05", {}, fmt=dtp.infer_format("date", "01 Oktober 2026")) == "05 November 2026"


def fake_ai_answer(original: bytes) -> dict:
    """What a correct AI answer looks like for sample_contract()."""
    spec = spec_for(original)
    for f in spec["fields"]:
        f["group"] = "client"
        f.setdefault("source", "")
        for o in f["occurrences"]:
            o.update(whole=False, remove=False)
    d = spec["deliverables"]
    return {"fields": spec["fields"], "deliverables": {"found": True, "table": 0, "template_row": 1, "sample_rows_after": 1,
            "columns": [dict(c, type=c.get("type", "text")) for c in d["columns"]]}}


def test_clean_spec_and_self_check_pass():
    from heroez_ops.contract_ai import clean_spec, self_check
    original = sample_contract()
    spec, notes = clean_spec(fake_ai_answer(original), original)
    assert notes == []
    assert spec["fields"][2]["format"] == {"prefix": "Rp", "suffix": "", "sep": "."}
    ok, msg = self_check(original, spec)
    assert ok, msg


def test_too_short_field_is_visible_in_review_preview():
    from heroez_ops.contract_ai import clean_spec, marked_preview, self_check
    original = sample_contract()
    raw = fake_ai_answer(original)
    raw["fields"][0]["occurrences"][0]["text"] = "PT Contoh"        # too short: leaves "Brand Tbk" behind
    spec, _ = clean_spec(raw, original)
    assert self_check(original, spec)[0]                              # positions still fine...
    assert "[[Nama klien]] Brand Tbk" in marked_preview(original, spec)   # ...but the reviewer sees the leftover


def test_clean_spec_drops_text_not_in_document():
    from heroez_ops.contract_ai import clean_spec
    original = sample_contract()
    raw = fake_ai_answer(original)
    raw["fields"].append({"key": "ghost", "label": "Hantu", "type": "text", "group": "other", "source": "",
                          "occurrences": [{"block": "B1", "text": "tidak ada di dokumen", "whole": False, "remove": False}]})
    spec, notes = clean_spec(raw, original)
    assert "ghost" not in {f["key"] for f in spec["fields"]}
    assert any("Hantu" in n for n in notes)
