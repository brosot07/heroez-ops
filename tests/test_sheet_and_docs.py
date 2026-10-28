import io
import json
import threading
import zipfile

import httpx
import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from heroez_ops import doc_import, sheet_sync, talents
from heroez_ops.db import Store
from heroez_ops.docs_ui import ContractsUI
from heroez_ops.kb import KB

OK = lambda url: (True, "")  # noqa: E731


def workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "KOL"
    ws.append(["Nama", "Handle IG", "Followers IG", "Niche", "Tipe", "Rate IG Reels", "No HP", "Handler"])
    ws.append(["Dewi Lestari", "@dewi.lestari", "120rb", "beauty", "eksklusif", 15000000, "0812-1111-2222", "budi"])
    ws.append(["Bima", "@bima.eats", "80rb", "fnb", "non-eksklusif", "", "", ""])
    j = wb.create_sheet("Jadwal")
    j.append(["Nama", "Mulai", "Selesai"])
    j.append(["Dewi Lestari", "2026-11-10", "2026-11-12"])
    f = wb.create_sheet("FAQ")
    f.append(["Pertanyaan", "Jawaban", "Jawaban EN", "Otomatis"])
    f.append(["Alamat kantor Heroez?\nkantornya dimana", "Jl. Contoh No. 1, Jakarta 🙏", "Jl. Contoh No. 1, Jakarta", "ya"])
    f.append(["Minimal budget?", "Mulai dari Rp25 juta per campaign.", "", "tidak"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_sheet_sync_reads_all_tabs():
    store = Store(":memory:")
    KB(store).set("sheet_url", "https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/edit#gid=0", "t")
    seen = []

    def handler(req):
        seen.append(str(req.url))
        return httpx.Response(200, content=workbook())
    st = sheet_sync.sync(store, http=httpx.Client(transport=httpx.MockTransport(handler)), check_host=OK)
    assert seen == ["https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/export?format=xlsx"]
    assert st["ok"] and (st["kol"], st["jadwal"], st["faq"]) == (2, 1, 2) and st["errors"] == []
    rows = {r["name"]: json.loads(r["data"]) for r in store.q("SELECT name, data FROM talents")}
    assert rows["Dewi Lestari"]["rates"] == {"Rate IG Reels": 15000000} and rows["Bima"]["type"] == "non-eksklusif"
    faqs = KB(store).faqs()
    assert faqs[0]["questions"] == ["Alamat kantor Heroez?", "kantornya dimana"] and faqs[0]["auto_send"]
    assert faqs[1]["auto_send"] is False
    assert KB(store).match_faq("Kantornya dimana?")["answer_id"].startswith("Jl. Contoh")


def test_sheet_not_shared_is_reported():
    store = Store(":memory:")
    KB(store).set("sheet_url", "https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/edit", "t")
    st = sheet_sync.sync(store, http=httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(302, headers={"location": "https://accounts.google.com"}))), check_host=OK)
    assert not st["ok"] and "belum dibagikan" in st["errors"][0]


def sample_pptx() -> bytes:
    """A two-slide deck built by hand (no python-pptx needed)."""
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    def slide(lines):
        body = "".join(f"<a:p><a:r><a:t>{t}</a:t></a:r></a:p>" for t in lines)
        return f'<?xml version="1.0"?><p:sld {ns}><p:cSld><p:spTree><p:sp><p:txBody>{body}</p:txBody></p:sp></p:spTree></p:cSld></p:sld>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("ppt/slides/slide1.xml", slide(["Dewi Lestari", "@dewi.lestari · 120K", "IG Reels Rp15.000.000"]))
        z.writestr("ppt/slides/slide2.xml", slide(["Sinta", "@sinta.glow · 300K", "IG Reels Rp25.000.000"]))
        z.writestr("ppt/media/image1.png", b"\x89PNG\r\n\x1a\n" + b"\0" * 100)
    return buf.getvalue()


def test_pptx_extraction_keeps_slide_order_and_images():
    blocks = doc_import.extract(sample_pptx(), "deck.pptx")
    texts = [b["text"] for b in blocks if b["type"] == "text"]
    assert texts[0].startswith("--- Slide 1 ---\nDewi Lestari") and "Sinta" in texts[1]
    assert [b["type"] for b in blocks].count("image") == 1
    with pytest.raises(doc_import.ImportError_):
        doc_import.extract(b"x", "deck.ppt")


def test_clean_and_roster_rows():
    raw = {"talents": [{"name": " Dewi Lestari ", "handles": ["IG @dewi.lestari"], "followers": ["IG 120K"], "niche": "beauty",
                        "type": "eksklusif", "rates": [{"label": "IG Reels", "amount_idr": 15000000, "from_price": False},
                                                       {"label": "x", "amount_idr": 0, "from_price": False}], "notes": ""},
                       {"name": "", "handles": [], "followers": [], "niche": "", "type": "", "rates": [], "notes": ""}],
           "document_notes": ["Harga belum termasuk pajak"]}
    res = doc_import.clean(raw)
    assert len(res["talents"]) == 1 and res["talents"][0]["rates"] == [{"label": "IG Reels", "amount_idr": 15000000, "from_price": False}]
    rows = doc_import.to_roster_rows(res["talents"])
    assert rows == [{"Nama": "Dewi Lestari", "Niche": "beauty", "Tipe": "eksklusif", "Handle IG": "@dewi.lestari",
                     "Followers IG": "120K", "Rate IG Reels": "15000000"}]


def test_merge_roster_keeps_private_columns():
    store = Store(":memory:")
    talents.import_roster(store, [{"Nama": "Dewi Lestari", "Niche": "beauty", "Rate IG Reels": "10jt", "No HP": "0812-1111-2222", "Handler": "budi"}], "t")
    added, updated, _ = talents.merge_roster(store, [{"Nama": "dewi lestari", "Rate IG Reels": "15000000", "Handle IG": "@dewi.lestari"},
                                                     {"Nama": "Sinta", "Niche": "beauty", "Rate IG Reels": "25000000"}], "t")
    assert (added, updated) == (1, 1)
    dewi = store.one("SELECT * FROM talents WHERE name='Dewi Lestari'")
    assert dewi["phone"] == "6281211112222" and dewi["handler"] == "budi"
    d = json.loads(dewi["data"])
    assert d["rates"] == {"Rate IG Reels": 15000000} and d["info"]["Handle IG"] == "@dewi.lestari" and d["info"]["Niche"] == "beauty"


@pytest.fixture
def client(cfg, tmp_path):
    import dataclasses
    c = dataclasses.replace(cfg, anthropic_api_key="key")
    ui = ContractsUI(tmp_path, Store(":memory:"), lambda: c)
    ui.doc_analyzer = lambda blocks, key, store: doc_import.clean({
        "talents": [{"name": "Dewi Lestari", "handles": ["IG @dewi.lestari"], "followers": ["IG 120K"], "niche": "beauty",
                     "type": "eksklusif", "rates": [{"label": "IG Reels", "amount_idr": 15000000, "from_price": False}], "notes": "cek angka"}],
        "document_notes": ["Harga belum termasuk pajak"]})
    app = FastAPI()
    app.include_router(ui.router(local=True))
    return ui, TestClient(app)


def test_document_review_and_approve_flow(client):
    ui, c = client
    r = c.post("/pengetahuan/dokumen", data={"csrf": ui.csrf}, files={"file": ("deck.pptx", sample_pptx())}, follow_redirects=False)
    assert r.status_code == 303
    did = int(r.headers["location"].rsplit("/", 1)[1])
    for t in threading.enumerate():
        if t.name == f"doc-import-{did}":
            t.join(10)
    page = c.get(f"/pengetahuan/dokumen/{did}").text
    assert "Dewi Lestari" in page and "IG Reels = 15000000" in page and "Harga belum termasuk pajak" in page
    assert ui.store.q("SELECT * FROM talents") == []                       # nothing until approved
    r = c.post(f"/pengetahuan/dokumen/{did}/setujui", data={
        "csrf": ui.csrf, "keep_0": "on", "name_0": "Dewi Lestari", "handles_0": "IG @dewi.lestari", "followers_0": "IG 120K",
        "niche_0": "beauty, skincare", "type_0": "eksklusif", "rates_0": "IG Reels = 16000000\nIG Story = 5000000"})
    assert "1 KOL baru" in r.text
    d = json.loads(ui.store.one("SELECT data FROM talents WHERE name='Dewi Lestari'")["data"])
    assert d["rates"] == {"Rate IG Reels": 16000000, "Rate IG Story": 5000000}      # the reviewer's corrections win
    assert d["info"]["Niche"] == "beauty, skincare"
    assert ui.store.one("SELECT status FROM doc_imports WHERE id=?", (did,))["status"] == "approved"


def test_unsupported_document_is_refused(client):
    ui, c = client
    r = c.post("/pengetahuan/dokumen", data={"csrf": ui.csrf}, files={"file": ("x.exe", b"MZ...")})
    assert "Format tidak didukung" in r.text
