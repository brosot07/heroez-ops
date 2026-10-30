import dataclasses
import io
import json
import threading
import zipfile

import docx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from heroez_ops import doctemplate as dtp
from heroez_ops import spk
from heroez_ops.db import Store
from heroez_ops.docs_ui import ContractsUI

BRIEF = """Halo kak, untuk campaign Glowskin Serum Launch (PT Glowskin Indonesia) periode 10-30 November 2026:
- Dewi Lestari: 2 video TikTok, tayang 20 Nov, fee 30jt
- Raka Pratama: 1 IG Reels tayang 22 Nov, fee 18jt
Pembayaran 50% di muka, pelunasan 14 hari setelah tayang."""

AI_JOB = {"client_name": "PT Glowskin Indonesia", "campaign": "Glowskin Serum Launch", "period_start": "2026-11-10",
          "period_end": "2026-11-30", "payment_terms": "50% di muka, pelunasan 14 hari setelah tayang",
          "rows": [{"talent": "Dewi Lestari", "platform": "TikTok", "deliverable": "Video 60 detik", "qty": 2,
                    "posting_date": "2026-11-20", "client_fee_idr": 30_000_000, "talent_fee_idr": 0, "notes": ""},
                   {"talent": "Raka Pratama", "platform": "Instagram", "deliverable": "Reels", "qty": 1,
                    "posting_date": "2026-11-22", "client_fee_idr": 18_000_000, "talent_fee_idr": 0, "notes": "usage 3 bulan"}],
          "uncertain": ["Fee talent tidak disebut di brief"]}


def sample_spk() -> bytes:
    """A small lawyer-style SPK from Heroez to a talent."""
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "No. 001/HRZ-SPK/X/2026"
    d.add_heading("SURAT PERINTAH KERJA", level=1)
    d.add_paragraph("PT Heroez Kreasi Indonesia memberikan perintah kerja kepada:")
    d.add_paragraph("Nama talent: Dewi Lestari")
    d.add_paragraph("Alamat: Jl. Kemang Raya 10, Jakarta Selatan")
    d.add_paragraph("Rekening: BCA 1234567890 a.n. Dewi Lestari")
    d.add_paragraph("untuk kampanye Glowskin Serum Launch dari brand PT Glowskin Indonesia, "
                    "periode 10 November 2026 sampai 30 November 2026.")
    t = d.add_table(rows=2, cols=5)
    for c, h in zip(t.rows[0].cells, ["No", "Platform", "Deliverable", "Tanggal tayang", "Fee"]):
        c.text = h
    for c, v in zip(t.rows[1].cells, ["1", "TikTok", "Video 60 detik", "20/11/2026", "Rp10.000.000"]):
        c.text = v
    d.add_paragraph("Total fee talent Rp10.000.000 (sepuluh juta Rupiah), dibayar 14 hari kerja setelah konten tayang.")
    d.add_paragraph("Pasal 1. Talent wajib mengikuti brief dan tidak mengunggah konten pesaing selama periode kampanye.")
    d.add_paragraph("Jakarta, 1 Oktober 2026")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def spk_ai_answer(original: bytes) -> dict:
    """What a correct template reading looks like for sample_spk()."""
    blocks = {dtp.run_text(b.paragraph): b.id for b in dtp.blocks(dtp.load(original))}
    b = lambda start: next(i for t, i in blocks.items() if t.startswith(start))  # noqa: E731
    occ = lambda block, text: [{"block": block, "text": text, "whole": False, "remove": False}]  # noqa: E731
    camp = b("untuk kampanye")
    fields = [
        ("spk_number", "Nomor SPK", "contract_number", "contract", occ(b("No. "), "001/HRZ-SPK/X/2026")),
        ("talent_name", "Nama talent", "text", "other", occ(b("Nama talent"), "Dewi Lestari")),
        ("talent_address", "Alamat talent", "text", "other", occ(b("Alamat"), "Jl. Kemang Raya 10, Jakarta Selatan")),
        ("talent_bank_account", "Rekening talent", "text", "other", occ(b("Rekening"), "BCA 1234567890 a.n. Dewi Lestari")),
        ("campaign_name", "Nama campaign", "text", "campaign", occ(camp, "Glowskin Serum Launch")),
        ("brand_name", "Nama brand", "text", "client", occ(camp, "PT Glowskin Indonesia")),
        ("period_start", "Mulai", "date", "campaign", occ(camp, "10 November 2026")),
        ("period_end", "Selesai", "date", "campaign", occ(camp, "30 November 2026")),
        ("fee_total", "Total fee talent", "money", "fee", occ(b("Total fee"), "Rp10.000.000")),
        ("fee_words", "Terbilang", "money_words", "fee", occ(b("Total fee"), "sepuluh juta Rupiah")),
        ("sign_date", "Tanggal SPK", "date", "contract", occ(b("Jakarta,"), "1 Oktober 2026")),
    ]
    return {"fields": [{"key": k, "label": lab, "type": ty, "group": g, "source": "fee_total" if ty == "money_words" else "",
                        "occurrences": o} for k, lab, ty, g, o in fields],
            "deliverables": {"found": True, "table": 0, "template_row": 1, "sample_rows_after": 0,
                             "columns": [{"col": 0, "key": "no", "label": "No", "type": "number"},
                                         {"col": 1, "key": "platform", "label": "Platform", "type": "text"},
                                         {"col": 2, "key": "deliverable", "label": "Deliverable", "type": "text"},
                                         {"col": 3, "key": "posting_date", "label": "Tanggal tayang", "type": "date"},
                                         {"col": 4, "key": "fee", "label": "Fee", "type": "money"}]}}


def wait(prefix: str):
    for t in threading.enumerate():
        if t.name.startswith(prefix):
            t.join(10)


@pytest.fixture
def app(cfg, tmp_path):
    from heroez_ops.contract_ai import clean_spec
    c = dataclasses.replace(cfg, anthropic_api_key="test-key", pkp=False)
    ui = ContractsUI(tmp_path, Store(":memory:"), lambda: c, analyzer=lambda data, key: clean_spec(spk_ai_answer(data), data))
    ui.spk_analyzer = lambda blocks, key, store: spk.clean(AI_JOB)
    a = FastAPI()
    a.include_router(ui.router(local=True))
    client = TestClient(a)
    r = client.post("/template/upload", data={"csrf": ui.csrf, "name": "SPK Talent v2026", "doc_type": "spk"},
                    files={"file": ("spk.docx", sample_spk())}, follow_redirects=False)
    tid = int(r.headers["location"].rsplit("/", 1)[1])
    wait(f"template-{tid}")
    assert client.post(f"/template/{tid}/setujui", data={"csrf": ui.csrf}, follow_redirects=False).status_code == 303
    return ui, client, tid


def new_job(ui, client) -> int:
    r = client.post("/spk/baru", data={"csrf": ui.csrf, "text": BRIEF, "title": "Glowskin Nov"},
                    files=[("files", ("brief.txt", BRIEF.encode()))], follow_redirects=False)
    jid = int(r.headers["location"].rsplit("/", 1)[1])
    wait(f"spk-{jid}")
    return jid


def test_drop_files_ai_fills_the_sow_table(app):
    ui, client, _ = app
    assert 'id="drop"' in client.get("/spk").text
    jid = new_job(ui, client)
    page = client.get(f"/spk/{jid}").text
    assert 'value="PT Glowskin Indonesia"' in page and 'value="Dewi Lestari"' in page and 'value="30.000.000"' in page
    assert "Fee talent tidak disebut di brief" in page                       # the AI's doubts are shown
    assert "fee talent kosong" in page
    j = ui.store.one("SELECT * FROM spk_jobs WHERE id=?", (jid,))
    assert json.loads(j["files"]) == ["brief.txt"] and (ui.home / "data" / "spk" / str(jid) / "_teks.txt").exists()


def test_commission_fills_talent_fees_then_one_spk_per_talent(app):
    ui, client, tid = app
    jid = new_job(ui, client)
    form = {"csrf": ui.csrf, "client_name": "PT Glowskin Indonesia", "campaign": "Glowskin Serum Launch",
            "period_start": "2026-11-10", "period_end": "2026-11-30", "payment_terms": "50/50", "commission_pct": "20",
            "r_talent_0": "Dewi Lestari", "r_platform_0": "TikTok", "r_deliverable_0": "Video 60 detik", "r_qty_0": "2",
            "r_posting_date_0": "2026-11-20", "r_client_fee_0": "30jt", "r_talent_fee_0": "", "r_notes_0": "",
            "r_talent_1": "Raka Pratama", "r_platform_1": "Instagram", "r_deliverable_1": "Reels", "r_qty_1": "1",
            "r_posting_date_1": "2026-11-22", "r_client_fee_1": "18.000.000", "r_talent_fee_1": "15jt", "r_notes_1": ""}
    r = client.post(f"/spk/{jid}/simpan", data=form, follow_redirects=False)
    assert "1%20fee%20talent%20diisi" in r.headers["location"] or "1 fee talent diisi" in r.headers["location"]
    rows = json.loads(ui.store.one("SELECT data FROM spk_jobs WHERE id=?", (jid,))["data"])["rows"]
    assert [x["talent_fee"] for x in rows] == [24_000_000, 15_000_000]       # typed fee kept, empty one = 80%
    # First time: address and bank are unknown → the batch skips and says why.
    r = client.post(f"/spk/{jid}/buat", data={"csrf": ui.csrf, "tid": tid}, follow_redirects=False)
    assert "Alamat%20talent" in r.headers["location"] or "Alamat talent" in r.headers["location"]
    assert ui.store.one("SELECT COUNT(*) n FROM contracts")["n"] == 0
    # Fill one by hand: the form is pre-filled from the table.
    form_html = client.get(f"/kontrak/baru?t={tid}&job={jid}&talent=Dewi%20Lestari").text
    assert 'value="Dewi Lestari"' in form_html and 'value="Glowskin Serum Launch"' in form_html
    assert 'name="r_fee_0" value="24.000.000"' in form_html and 'name="job" value="%d"' % jid in form_html
    number = __import__("re").search(r'name="f_spk_number" value="([^"]+)"', form_html).group(1)
    assert number.startswith("001/HRZ-SPK/")
    fields = {k: v for k, v in __import__("re").findall(r'name="(f_[a-z_]+)" value="([^"]*)"', form_html)}
    fields.update({"f_talent_address": "Jl. Melati 5, Jakarta", "f_talent_bank_account": "BCA 111 a.n. Dewi Lestari"})
    rows_f = dict(__import__("re").findall(r'name="(r_[a-z_]+_\d+)" value="([^"]*)"', form_html))
    r = client.post("/kontrak/simpan", data={"csrf": ui.csrf, "tid": tid, "cid": 0, "job": jid, "talent": "Dewi Lestari",
                                             "deal_id": 0, **fields, **rows_f}, follow_redirects=False)
    assert r.headers["location"] == f"/spk/{jid}"
    doc = ui.store.one("SELECT * FROM contracts WHERE talent='Dewi Lestari'")
    assert doc["doc_type"] == "spk" and doc["batch_id"] == jid and doc["number"] == number and doc["total"] == 24_000_000
    body = "\n".join(p.text for p in dtp.load((ui.cdir / f"{doc['id']}.docx").read_bytes()).paragraphs)
    assert "Nama talent: Dewi Lestari" in body and "Jl. Melati 5, Jakarta" in body
    assert "Rp24.000.000 (dua puluh empat juta Rupiah)" in body and "Pasal 1. Talent wajib" in body
    # A second job for the same talent: details remembered, so "make all" works in one click.
    jid2 = new_job(ui, client)
    client.post(f"/spk/{jid2}/simpan", data={**form, "commission_pct": "20", "r_talent_1": "", "r_deliverable_1": ""})
    r = client.post(f"/spk/{jid2}/buat", data={"csrf": ui.csrf, "tid": tid}, follow_redirects=False)
    assert "1%20SPK%20dibuat" in r.headers["location"] or "1 SPK dibuat" in r.headers["location"]
    doc2 = ui.store.one("SELECT * FROM contracts WHERE batch_id=?", (jid2,))
    assert doc2["number"].startswith("002/HRZ-SPK/")
    body2 = "\n".join(p.text for p in dtp.load((ui.cdir / f"{doc2['id']}.docx").read_bytes()).paragraphs)
    assert "Jl. Melati 5, Jakarta" in body2 and "BCA 111 a.n. Dewi Lestari" in body2
    # SPKs never touch the deal pipeline; the ZIP holds the job's documents
    assert ui.store.one("SELECT COUNT(*) n FROM deals")["n"] == 0
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/spk/{jid}/zip").content))
    assert [n for n in z.namelist()] == [f"{number.replace('/', '-')} Dewi Lestari.docx"]


def test_field_meanings():
    m = spk.field_meaning
    assert m("talent_name", "Nama talent", "other") == "talent"
    assert m("talent_address", "Alamat talent", "other") is None
    assert m("talent_phone", "No. HP talent", "other") == "talent_phone"
    assert m("brand_name", "Nama brand", "client") == "client"
    assert m("client_address", "Alamat klien", "client") is None
    assert m("periode", "Periode kampanye", "campaign") == "period"
    assert m("campaign_name", "Nama campaign", "campaign") == "campaign"
    assert m("pihak_pertama", "Pihak Pertama", "other") is None              # which party is Heroez differs per template
    assert m("tanggal_mulai", "Tanggal mulai", "campaign", "date") == "period_start"
    assert m("sign_date", "Tanggal SPK", "contract", "date") is None          # → today
    assert spk.column_meaning("fee", "Fee") == "fee" and spk.column_meaning("tgl", "Tanggal tayang") == "posting_date"


def test_without_ai_key_the_table_is_filled_by_hand(app):
    ui, client, _ = app
    ui.cfg = (lambda c: (lambda: c))(dataclasses.replace(ui.cfg(), anthropic_api_key=""))
    jid = new_job(ui, client)
    j = ui.store.one("SELECT * FROM spk_jobs WHERE id=?", (jid,))
    assert j["status"] == "review" and "isi tabel di bawah secara manual" in j["notes"]
