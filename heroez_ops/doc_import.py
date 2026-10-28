"""Read any document (PowerPoint, PDF, Word, Excel, images) into a KOL table for review.

The AI copies what the document says into a table; a supervisor checks and approves it
before it touches the roster. Nothing reaches clients from this step.
"""
from __future__ import annotations

import base64
import datetime as dt
import io
import json
import logging
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from .db import Store
from .safety import UnsafeFile, check_zip
from .terbilang import rupiah

log = logging.getLogger(__name__)

PROMPT = (Path(__file__).resolve().parent / "prompts" / "kol_from_document_v1.md").read_text(encoding="utf-8")
MODEL = "claude-sonnet-5-5"
PRICES = {"claude-sonnet-5-5": (2.0, 10.0, 0.20)}
MAX_IMAGES, MAX_IMAGE_BYTES, MAX_TOTAL_BYTES = 20, 3 * 1024 * 1024, 20 * 1024 * 1024
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".webp": "image/webp"}

SCHEMA = {
    "type": "object",
    "properties": {
        "talents": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "name": {"type": "string"},
                "handles": {"type": "array", "items": {"type": "string"}},
                "followers": {"type": "array", "items": {"type": "string"}},
                "niche": {"type": "string"},
                "type": {"type": "string", "enum": ["eksklusif", "non-eksklusif", "nano", ""]},
                "rates": {"type": "array", "items": {
                    "type": "object",
                    "properties": {"label": {"type": "string"}, "amount_idr": {"type": "integer"}, "from_price": {"type": "boolean"}},
                    "required": ["label", "amount_idr", "from_price"], "additionalProperties": False}},
                "notes": {"type": "string"},
            },
            "required": ["name", "handles", "followers", "niche", "type", "rates", "notes"],
            "additionalProperties": False}},
        "document_notes": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["talents", "document_notes"],
    "additionalProperties": False,
}


class ImportError_(Exception):
    pass


# ------------------------------------------------------------------ extraction (no AI)

def _pptx_blocks(data: bytes) -> list[dict]:
    """Slide texts in order plus the pictures embedded in the deck."""
    check_zip(data)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    blocks, images, total = [], [], 0
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                        key=lambda n: int(re.search(r"(\d+)", n).group(1)))
        for n in slides:
            root = ET.fromstring(z.read(n))
            texts = []
            for p in root.iter("{%s}p" % ns["a"]):
                t = "".join(r.text or "" for r in p.iter("{%s}t" % ns["a"])).strip()
                if t:
                    texts.append(t)
            blocks.append({"type": "text", "text": f"--- Slide {slides.index(n) + 1} ---\n" + "\n".join(texts)})
        for n in z.namelist():
            ext = Path(n).suffix.lower()
            if n.startswith("ppt/media/") and ext in IMAGE_TYPES and len(images) < MAX_IMAGES:
                b = z.read(n)
                if len(b) <= MAX_IMAGE_BYTES and total + len(b) <= MAX_TOTAL_BYTES:
                    images.append({"type": "image", "source": {"type": "base64", "media_type": IMAGE_TYPES[ext],
                                                               "data": base64.b64encode(b).decode()}})
                    total += len(b)
    return blocks + images


def _docx_text(data: bytes) -> str:
    check_zip(data)
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    from .sheet_sync import read_workbook
    check_zip(data)
    out = []
    for title, rows in read_workbook(data).items():
        out.append(f"--- Sheet {title} ---")
        for r in rows[:500]:
            out.append(" | ".join(f"{k}: {v}" for k, v in r.items() if v))
    return "\n".join(out)


def extract(data: bytes, filename: str) -> list[dict]:
    """Content blocks for the AI call: text, images or a PDF document."""
    ext = Path(filename).suffix.lower()
    if len(data) > MAX_TOTAL_BYTES:
        raise ImportError_("File terlalu besar (maks. 20 MB).")
    if ext == ".pptx":
        return _pptx_blocks(data)
    if ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise ImportError_("File bukan PDF.")
        return [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                 "data": base64.b64encode(data).decode()}}]
    if ext == ".docx":
        return [{"type": "text", "text": _docx_text(data)}]
    if ext in (".xlsx",):
        return [{"type": "text", "text": _xlsx_text(data)}]
    if ext in (".csv", ".txt", ".md"):
        return [{"type": "text", "text": data.decode("utf-8", errors="replace")[:200_000]}]
    if ext in IMAGE_TYPES:
        if len(data) > MAX_IMAGE_BYTES:
            raise ImportError_("Gambar terlalu besar (maks. 3 MB).")
        return [{"type": "image", "source": {"type": "base64", "media_type": IMAGE_TYPES[ext],
                                             "data": base64.b64encode(data).decode()}}]
    if ext == ".ppt":
        raise ImportError_("Format .ppt lama tidak didukung: simpan ulang sebagai .pptx atau PDF.")
    raise ImportError_("Format tidak didukung. Pakai .pptx, .pdf, .docx, .xlsx, .csv atau gambar.")


# ------------------------------------------------------------------ AI call

def analyze(blocks: list[dict], api_key: str, store: Store, client=None) -> dict:
    import anthropic
    client = client or anthropic.Anthropic(api_key=api_key, timeout=600.0, max_retries=2)
    if not any(b["type"] != "text" or b.get("text", "").strip() for b in blocks):
        return {"talents": [], "document_notes": ["Dokumen kosong atau tidak terbaca."]}
    content = blocks + [{"type": "text", "text": "Salin semua KOL dan harga dari dokumen di atas ke dalam tabel."}]
    try:
        resp = client.beta.messages.create(
            model=MODEL, max_tokens=16000,
            betas=["server-side-fallback-2026-07-01"], fallbacks="default",
            output_config={"effort": "medium", "format": {"type": "json_schema", "schema": SCHEMA}},
            system=PROMPT,
            messages=[{"role": "user", "content": content}],
        )
    except anthropic.AuthenticationError as e:
        raise ImportError_("Claude API key ditolak. Cek di halaman Pengaturan.") from e
    except (anthropic.APIConnectionError, anthropic.RateLimitError, anthropic.APIStatusError) as e:
        raise ImportError_(f"Claude tidak bisa dihubungi ({type(e).__name__}). Coba lagi.") from e
    _meter(store, resp.usage)
    if resp.stop_reason != "end_turn":
        raise ImportError_(f"AI berhenti tanpa hasil ({resp.stop_reason}).")
    data = json.loads(next(b.text for b in resp.content if b.type == "text"))
    return clean(data)


def clean(data: dict) -> dict:
    talents = []
    for t in data.get("talents", []):
        name = (t.get("name") or "").strip()[:80]
        if not name:
            continue
        rates = []
        for r in t.get("rates", []):
            try:
                amt = int(r.get("amount_idr", 0))
            except (TypeError, ValueError):
                amt = 0
            if amt > 0 and (r.get("label") or "").strip():
                rates.append({"label": r["label"].strip()[:60], "amount_idr": amt, "from_price": bool(r.get("from_price"))})
        talents.append({"name": name, "handles": [h.strip()[:60] for h in t.get("handles", []) if h.strip()][:6],
                        "followers": [f.strip()[:40] for f in t.get("followers", []) if f.strip()][:6],
                        "niche": (t.get("niche") or "").strip()[:120], "type": t.get("type") or "",
                        "rates": rates[:12], "notes": (t.get("notes") or "").strip()[:300]})
    return {"talents": talents, "document_notes": [n.strip()[:200] for n in data.get("document_notes", []) if n.strip()][:20]}


def _meter(store: Store, usage) -> None:
    pin, pout, pcache = PRICES[MODEL]
    inp, out = usage.input_tokens or 0, usage.output_tokens or 0
    cread = getattr(usage, "cache_read_input_tokens", 0) or 0
    cwrite = getattr(usage, "cache_creation_input_tokens", 0) or 0
    usd = (inp * pin + cwrite * pin * 1.25 + cread * pcache + out * pout) / 1_000_000
    store.x("INSERT INTO ai_usage(day,model,calls,input_tokens,output_tokens,cache_read_tokens,usd) VALUES(?,?,1,?,?,?,?) "
            "ON CONFLICT(day,model) DO UPDATE SET calls=calls+1, input_tokens=input_tokens+excluded.input_tokens, "
            "output_tokens=output_tokens+excluded.output_tokens, cache_read_tokens=cache_read_tokens+excluded.cache_read_tokens, "
            "usd=usd+excluded.usd", (dt.date.today().isoformat(), MODEL, inp + cwrite, out, cread, usd))


# ------------------------------------------------------------------ to roster rows

def to_roster_rows(talents: list[dict]) -> list[dict[str, str]]:
    """The review table as roster rows (same columns as the KOL sheet), ready for merge_roster."""
    rows = []
    for t in talents:
        row = {"Nama": t["name"], "Niche": t.get("niche", ""), "Tipe": t.get("type", "")}
        for h in t.get("handles", []):
            plat, _, handle = h.partition(" ")
            key = f"Handle {plat}" if handle else "Handle"
            row[key] = handle or plat
        for f in t.get("followers", []):
            plat, _, n = f.partition(" ")
            row[f"Followers {plat}" if n else "Followers"] = n or plat
        for r in t.get("rates", []):
            label = r["label"] if not r.get("from_price") else f"{r['label']} (mulai dari)"
            row[f"Rate {label}"] = str(r["amount_idr"])
        rows.append(row)
    return rows


def preview_rate(r: dict) -> str:
    return ("mulai dari " if r.get("from_price") else "") + rupiah(r["amount_idr"])
