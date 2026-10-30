"""Fill approved Word (.docx) contract templates without changing their wording.

A template is the lawyer-approved .docx with some text marked as fields
(client name, fee, SOW rows...). Marking replaces that text with {{key}}
placeholders, keeping the run formatting around it. Rendering puts values back
in. Approved clauses are never rewritten.

Block ids (B0, B1, ...) number every paragraph in reading order: body text,
table cells (top-level tables), then headers and footers.
"""
from __future__ import annotations

import datetime as dt
import io
import re
from copy import deepcopy
from dataclasses import dataclass

import docx
from docx.table import _Row
from docx.text.paragraph import Paragraph

from .terbilang import rupiah_words

PLACEHOLDER = re.compile(r"\{\{\s*([a-z0-9_]+(?:\.[a-z0-9_]+)?)\s*\}\}")
BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober",
         "November", "Desember"]


@dataclass
class Block:
    id: str
    paragraph: Paragraph
    where: str = ""        # "T2 R1 C3" for table cells, "header"/"footer", "" for body


def load(data: bytes) -> docx.document.Document:
    return docx.Document(io.BytesIO(data))


def save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def blocks(doc) -> list[Block]:
    out: list[Block] = []
    body = doc.element.body
    tables = {t._tbl: i for i, t in enumerate(doc.tables)}
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            out.append(Block(f"B{len(out)}", Paragraph(child, doc._body)))
        elif child.tag.endswith("}tbl") and child in tables:
            ti = tables[child]
            for ri, row in enumerate(doc.tables[ti].rows):
                seen = set()
                for ci, cell in enumerate(row.cells):
                    if id(cell._tc) in seen:   # merged cells repeat
                        continue
                    seen.add(id(cell._tc))
                    for p in cell.paragraphs:
                        out.append(Block(f"B{len(out)}", p, f"T{ti} R{ri} C{ci}"))
    for si, section in enumerate(doc.sections):
        for kind, part in (("header", section.header), ("footer", section.footer)):
            if part.is_linked_to_previous and si > 0:
                continue
            for p in part.paragraphs:
                out.append(Block(f"B{len(out)}", p, kind))
    return out


def run_text(p: Paragraph) -> str:
    return "".join(r.text for r in p.runs)


def outline(doc) -> str:
    """Numbered text of the document, for the AI and the review screen."""
    lines = []
    for b in blocks(doc):
        t = run_text(b.paragraph)
        if t.strip():
            lines.append(f"{b.id}{' [' + b.where + ']' if b.where else ''}: {t}")
    return "\n".join(lines)


def replace_in_paragraph(p: Paragraph, old: str, new: str) -> bool:
    """Replace the first `old` in a paragraph, keeping the formatting of the run where it starts."""
    if not old:
        return False
    runs = p.runs
    full = "".join(r.text for r in runs)
    i = full.find(old)
    if i < 0:
        return False
    j = i + len(old)
    pos, first, last = 0, None, None
    for k, r in enumerate(runs):
        n = len(r.text)
        if first is None and i < pos + n:
            first, f_off = k, i - pos
        if j <= pos + n:
            last, l_off = k, j - pos
            break
        pos += n
    if first == last:
        r = runs[first]
        r.text = r.text[:f_off] + new + r.text[l_off:]
    else:
        runs[first].text = runs[first].text[:f_off] + new
        for k in range(first + 1, last):
            runs[k].text = ""
        runs[last].text = runs[last].text[l_off:]
    return True


def replace_paragraph_text(p: Paragraph, new: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(new)
        return
    runs[0].text = new
    for r in runs[1:]:
        r.text = ""


def delete_paragraph(p: Paragraph) -> None:
    el = p._element
    el.getparent().remove(el)


# -- marking a template ---------------------------------------------------------

def mark(original: bytes, spec: dict) -> tuple[bytes, list[str]]:
    """Insert {{key}} placeholders described by spec. Returns (template bytes, problems)."""
    doc = load(original)
    by_id = {b.id: b for b in blocks(doc)}
    problems: list[str] = []
    to_delete: list[Paragraph] = []
    for f in spec.get("fields", []):
        for occ in f.get("occurrences", []):
            b = by_id.get(occ.get("block", ""))
            if not b:
                problems.append(f"{f['label']}: blok {occ.get('block')} tidak ada")
                continue
            if occ.get("remove"):
                to_delete.append(b.paragraph)
            elif occ.get("whole"):
                replace_paragraph_text(b.paragraph, "{{" + f["key"] + "}}")
            elif not replace_in_paragraph(b.paragraph, occ.get("text", ""), "{{" + f["key"] + "}}"):
                problems.append(f"{f['label']}: teks “{occ.get('text', '')[:40]}” tidak ditemukan di {b.id}")
    d = spec.get("deliverables")
    if d:
        try:
            row = doc.tables[d["table"]].rows[d["template_row"]]
            cells = []
            for c in row.cells:
                if c._tc not in [x._tc for x in cells]:
                    cells.append(c)
            for col in d["columns"]:
                replace_paragraph_text(cells[col["col"]].paragraphs[0], "{{row." + col["key"] + "}}")
                for extra in cells[col["col"]].paragraphs[1:]:
                    delete_paragraph(extra)
            # Other existing data rows below the template row are sample rows: drop them.
            rows = doc.tables[d["table"]].rows
            for extra in list(rows)[d["template_row"] + 1: d["template_row"] + 1 + int(d.get("sample_rows_after", 0))]:
                extra._tr.getparent().remove(extra._tr)
        except (IndexError, KeyError) as ex:
            problems.append(f"Tabel SOW tidak bisa ditandai: {ex}")
    for p in to_delete:
        delete_paragraph(p)
    return save(doc), problems


# -- rendering ------------------------------------------------------------------

MONTHS_EN = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
             "November", "December"]
_MONEY_RE = re.compile(r"^(?P<prefix>\D*?)(?P<num>\d{1,3}(?:[.,]\d{3})+|\d+)(?P<dec>[.,]00)?(?P<suffix>\D*)$")


def infer_format(ftype: str, example: str) -> dict:
    """Learn how the approved template writes this value, so output matches its style."""
    ex = (example or "").strip()
    if ftype == "money":
        m = _MONEY_RE.match(ex)
        if m:
            sep = "," if "," in m.group("num") and "." not in m.group("num") else "."
            return {"prefix": m.group("prefix"), "suffix": (m.group("dec") or "") + m.group("suffix"), "sep": sep}
        return {"prefix": "Rp", "suffix": "", "sep": "."}
    if ftype == "money_words":
        letters = [c for c in ex if c.isalpha()]
        case = "upper" if letters and all(c.isupper() for c in letters) else "title" if ex.istitle() else "lower"
        return {"case": case, "rupiah": "rupiah" in ex.lower()}
    if ftype == "date":
        if re.match(r"^\d{1,2}/\d{1,2}/\d{4}$", ex):
            return {"style": "dd/mm/yyyy"}
        if re.match(r"^\d{1,2}-\d{1,2}-\d{4}$", ex):
            return {"style": "dd-mm-yyyy"}
        if any(m in ex for m in MONTHS_EN):
            return {"style": "en"}
        return {"style": "id", "pad": bool(re.match(r"^0\d ", ex))}
    return {}


def tanggal(d: dt.date | str, style: str = "id", pad: bool = False) -> str:
    if isinstance(d, str):
        d = dt.date.fromisoformat(d)
    if style == "dd/mm/yyyy":
        return d.strftime("%d/%m/%Y")
    if style == "dd-mm-yyyy":
        return d.strftime("%d-%m-%Y")
    day = f"{d.day:02d}" if pad else str(d.day)
    return f"{day} {(MONTHS_EN if style == 'en' else BULAN)[d.month - 1]} {d.year}"


def format_value(ftype: str, value, values: dict, source: str | None = None, fmt: dict | None = None) -> str:
    fmt = fmt or {}
    if ftype == "money":
        n = f"{int(value or 0):,}".replace(",", fmt.get("sep", "."))
        return f"{fmt.get('prefix', 'Rp')}{n}{fmt.get('suffix', '')}"
    if ftype == "money_words":
        words = rupiah_words(int(values.get(source or "", 0) or 0))
        if not fmt.get("rupiah", True):
            words = words.removesuffix(" Rupiah")
        case = fmt.get("case", "lower")
        return words.upper() if case == "upper" else words.title() if case == "title" else words
    if ftype == "date":
        return tanggal(value, fmt.get("style", "id"), fmt.get("pad", False)) if value else ""
    if ftype == "number":
        return f"{int(value):,}".replace(",", ".") if str(value).strip() else ""
    return "" if value is None else str(value)


def render(template: bytes, spec: dict, values: dict, rows: list[dict] | None = None) -> bytes:
    """Fill a marked template. values: raw values by field key (money as int, dates ISO)."""
    doc = load(template)
    types = {f["key"]: f for f in spec.get("fields", [])}
    text_values = {k: format_value(f.get("type", "text"), values.get(k), values, f.get("source"), f.get("format"))
                   for k, f in types.items()}

    d = spec.get("deliverables")
    if d:
        table = doc.tables[d["table"]]
        tmpl_tr = table.rows[d["template_row"]]._tr
        money_cols = {c["key"]: c.get("format") for c in d["columns"] if c.get("type") == "money"}
        date_cols = {c["key"]: c.get("format") for c in d["columns"] if c.get("type") == "date"}

        def cell_value(key: str, vals: dict) -> str | None:
            if not key.startswith("row."):
                return None
            k = key[4:]
            if k in money_cols:
                return format_value("money", vals.get(k, 0), vals, fmt=money_cols[k])
            if k in date_cols and vals.get(k):
                try:
                    return format_value("date", vals[k], vals, fmt=date_cols[k])    # same style as the template
                except ValueError:
                    return str(vals[k])                                              # free text like "minggu ke-3"
            return str(vals.get(k, ""))

        for n, row in enumerate(rows or [], 1):
            new_tr = deepcopy(tmpl_tr)
            tmpl_tr.addprevious(new_tr)
            vals = {**row, "no": str(n)}
            for cell in _Row(new_tr, table).cells:
                for p in cell.paragraphs:
                    _fill(p, lambda key, vals=vals: cell_value(key, vals))
        tmpl_tr.getparent().remove(tmpl_tr)

    for b in blocks(doc):
        _fill(b.paragraph, lambda key: text_values.get(key))
    return save(doc)


def _fill(p: Paragraph, lookup) -> None:
    for _ in range(50):   # bounded: a value could itself look like a placeholder
        m = PLACEHOLDER.search(run_text(p))
        if not m:
            return
        value = lookup(m.group(1))
        if value is None:
            value = ""
        if not replace_in_paragraph(p, m.group(0), value):
            return


def placeholders(template: bytes) -> set[str]:
    doc = load(template)
    found = set()
    for b in blocks(doc):
        found.update(PLACEHOLDER.findall(run_text(b.paragraph)))
    return found
